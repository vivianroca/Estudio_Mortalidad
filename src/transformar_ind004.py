"""IND-004, 2009–2024: reproducir las categorías de causa publicadas."""

import csv
import hashlib
import re
from collections import Counter
from pathlib import Path

import pandas as pd
import xlrd
import openpyxl
from openpyxl.styles import Alignment, Font

from transformar_ind002 import DEPARTAMENTOS, entero, normalizar, unico

RAIZ = Path(__file__).resolve().parents[1]
ANIOS = tuple(range(2009, 2025))
CATALOGO = RAIZ / 'reports/estructura/catalogo_indicadores.csv'
SALIDA = RAIZ / 'data/processed/ind004_2009_2024.csv'
REPORTE = RAIZ / 'reports/estructura/validacion_ind004_2009_2024.txt'
FRECUENCIAS = RAIZ / 'reports/estructura/frecuencia_causas_ind004_2009_2024.xlsx'
COLUMNAS = ['anio','depto_residencia','causa_muerte','sexo','conteo','tiene_cie10']
SEXOS = ('Hombres','Mujeres')
NACIONAL = 'Todos los departamentos'
TODAS = 'Todas las causas'
ESPERADAS = {normalizar(d) for d in (*DEPARTAMENTOS,'Extranjero','Ignorado')}


def leer_hoja(anio, nombre):
    rutas = [RAIZ / 'data/raw_data' / f'{anio}{ext}' for ext in ('.xls','.xlsx')
             if (RAIZ / 'data/raw_data' / f'{anio}{ext}').is_file()]
    ruta = unico(rutas,f'Archivo del año {anio}')
    huella = hashlib.sha256(ruta.read_bytes()).hexdigest()
    if ruta.suffix == '.xls':
        libro = xlrd.open_workbook(str(ruta),on_demand=True)
        try:
            hoja = libro.sheet_by_name(nombre)
            matriz = [hoja.row_values(r) for r in range(hoja.nrows)]
            dimensiones = (hoja.nrows,hoja.ncols)
        finally:
            libro.release_resources()
    else:
        libro = openpyxl.load_workbook(ruta,read_only=True,data_only=False,keep_links=False)
        try:
            hoja = libro[nombre]
            dimensiones = (hoja.max_row,hoja.max_column)
            matriz = [[v if v is not None else '' for v in fila]
                      for fila in hoja.iter_rows(values_only=True)]
        finally:
            libro.close()
    return ruta,huella,matriz,dimensiones


def transformar(anio, nombre, informe):
    ruta,huella,matriz,dimensiones = leer_hoja(anio,nombre)
    informe.extend([f'\nAÑO: {anio}',f'Archivo: {ruta.name}; hoja: {nombre}',
                    f'Dimensiones: {dimensiones}; SHA-256 original: {huella}'])
    etiquetas = ('Departamento de residencia','Causa de muerte','Total',*SEXOS)
    r_enc = unico([r for r,fila in enumerate(matriz) if all(e in fila for e in etiquetas)],'Encabezado')
    cols = {e:unico([c for c,v in enumerate(matriz[r_enc]) if v==e],e) for e in etiquetas}
    nacional = unico(list(dict.fromkeys(fila[cols[etiquetas[0]]] for fila in matriz[r_enc+1:]
                     if normalizar(fila[cols[etiquetas[0]]]) in
                     {'todos los departamentos','total pais'})), 'Agregado nacional')
    cie = [c for c,v in enumerate(matriz[r_enc]) if normalizar(v)=='codigo cie-10']
    if len(cie)>1:
        raise ValueError('Más de una columna de código CIE-10')
    filas,errores,advertencias,notas = [],[],[],[]
    for r in range(r_enc+1,len(matriz)):
        fila = matriz[r]
        if not any(v!='' for v in fila):
            continue
        d,c = fila[cols[etiquetas[0]]],fila[cols[etiquetas[1]]]
        if isinstance(d,str) and normalizar(d).startswith(('nota:','notas:','fuente:')):
            notas.append(f'Fila {r+1}: {d}')
            if any(v!='' for i,v in enumerate(fila) if i!=cols[etiquetas[0]]):
                errores.append(f'Fila {r+1}: nota con celdas adicionales')
            continue
        if not isinstance(d,str) or not d.strip() or not isinstance(c,str) or not c.strip():
            errores.append(f'Fila {r+1}: residencia/causa vacía o no textual')
            continue
        if d!=nacional and normalizar(d) not in ESPERADAS:
            errores.append(f'Fila {r+1}: residencia desconocida {d!r}')
        if any(v!='' for i,v in enumerate(fila) if i not in set(cols.values())|set(cie)):
            errores.append(f'Fila {r+1}: celda fuera de las columnas identificadas')
        v={}
        for sexo in ('Total',*SEXOS):
            try:
                v[sexo]=entero(fila[cols[sexo]],f'fila {r+1}, columna {cols[sexo]+1}')
            except ValueError as error:
                errores.append(str(error))
        codigo = str(fila[cie[0]]).strip() if cie else ''
        flag = int(bool(re.fullmatch(r'[A-Z]\d{2}(?:\.\d+)?(?:\s*[-–]\s*[A-Z]\d{2}(?:\.\d+)?)?',codigo)))
        # Identidad interna por fila original: nunca agrupar por descripción o código.
        filas.append(dict(fila=r+1,depto=d,causa=c,valores=v,flag=flag))
    residencias=list(dict.fromkeys(f['depto'] for f in filas if f['depto']!=nacional))
    normalizadas=Counter(normalizar(d) for d in residencias)
    if set(normalizadas)!=ESPERADAS or any(n>1 for n in normalizadas.values()):
        errores.append('Residencias faltantes, desconocidas o inconsistentes')
    totales={}
    for d in [nacional,*residencias]:
        totales[d]=unico([f for f in filas if f['depto']==d and f['causa']==TODAS],f'Total de {d}')
    informe.extend([f'Encabezado: fila {r_enc+1}; columnas: {cols}; columna CIE presente: {bool(cie)}.',
                    'tiene_cie10=1 si la fila fuente tiene un código o rango CIE-10 explícito; 0 si no tiene código, no hay columna o es una etiqueta residual.',
                    'El flag registra el código publicado; no certifica pertenencia a un diccionario externo ni corrige la descripción.',
                    'No se valida igualdad de descripciones frente a códigos distintos. Se conservan todas las filas originales, sin deduplicar ni fusionar.',
                    'Agregado nacional y Todas las causas se excluyen del CSV; Otras causas, Extranjero e Ignorado se conservan.',
                    'Notas/fuentes: '+('; '.join(notas) or 'ninguna.'),
                    'No se imputan guiones, vacíos ni valores no numéricos.'])
    if errores:
        informe.extend('ERROR: '+e for e in errores)
        return None,None
    analiticas=[f for f in filas if f['depto']!=nacional and f['causa']!=TODAS]
    ancho=pd.DataFrame([dict(_fila=f['fila'],depto_residencia=f['depto'],causa_muerte=f['causa'],
                             tiene_cie10=f['flag'],**{s:f['valores'][s] for s in SEXOS}) for f in analiticas])
    largo=ancho.melt(id_vars=['_fila','depto_residencia','causa_muerte','tiene_cie10'],
                      value_vars=list(SEXOS),var_name='sexo',value_name='conteo')
    largo.insert(0,'anio',anio)
    fuente={(f['fila'],s):(f['depto'],f['causa'],f['valores'][s],f['flag']) for f in analiticas for s in SEXOS}
    salida={(r._0,r.sexo):(r.depto_residencia,r.causa_muerte,r.conteo,r.tiene_cie10)
            for r in largo[['_fila','sexo','depto_residencia','causa_muerte','conteo','tiene_cie10']].itertuples(index=False)}
    if salida!=fuente or len(largo)!=len(fuente):
        errores.append('Intersecciones incorrectas o pérdida de filas originales')
    for f in filas:
        if f['valores']['Hombres']+f['valores']['Mujeres']!=f['valores']['Total']:
            errores.append(f'Fila {f["fila"]}: Hombres+Mujeres no coincide con Total')
    informe.append('\nTOTALES POR RESIDENCIA Y SEXO:')
    for d in [nacional,*residencias]:
        partes=[f for f in filas if f['depto']==d and f['causa']!=TODAS]
        for sexo in ('Total',*SEXOS):
            calculado=sum(f['valores'][sexo] for f in partes)
            publicado=totales[d]['valores'][sexo]
            informe.append(f'{d}/{sexo}: publicado={publicado}; suma de causas={calculado}; diferencia={calculado-publicado}.')
            if calculado!=publicado:
                texto=f'{d}/{sexo}: publicado={publicado}, suma={calculado}, diferencia={calculado-publicado}; sin corrección.'
                # El bloque nacional no forma parte de los registros exportados.
                (advertencias if d==nacional else errores).append(texto)
    informe.append('\nTOTALES DEL CSV FRENTE AL TOTAL NACIONAL:')
    for sexo in ('Total',*SEXOS):
        calculado=int(largo.conteo.sum() if sexo=='Total' else largo.loc[largo.sexo==sexo,'conteo'].sum())
        publicado=totales[nacional]['valores'][sexo]
        informe.append(f'{sexo}: publicado={publicado}; transformado={calculado}; diferencia={calculado-publicado}.')
        if calculado!=publicado:
            errores.append(f'Total del CSV no coincide para {sexo}')
        if sum(totales[d]['valores'][sexo] for d in residencias)!=publicado:
            errores.append(f'Márgenes originales de residencia no coinciden para {sexo}')
    if hashlib.sha256(ruta.read_bytes()).hexdigest()!=huella:
        errores.append('El Excel original cambió durante la ejecución')
    informe.extend('ADVERTENCIA: '+a for a in advertencias)
    informe.append('Las listas de causas son selectivas; no se completan categorías ausentes con cero ni se asume composición constante de Otras causas.')
    if errores:
        informe.extend('ERROR: '+e for e in errores)
        return None,None
    largo=largo[COLUMNAS].astype({'conteo':'int64','tiene_cie10':'int64'})
    informe.extend([f'Intersecciones verificadas por fila fuente y sexo: {len(largo)}.',
                    'Guiones, vacíos y no numéricos en el bloque y márgenes: 0.',
                    f'RESULTADO {anio}: REPRODUCCIÓN VERIFICADA'+(' CON ADVERTENCIAS.' if advertencias else '.'),
                    'EJEMPLOS:',largo.head(4).to_csv(index=False).strip()])
    estructura=dict(dimensiones=dimensiones,encabezados=matriz[r_enc],columnas=cols,
                    residencias=residencias,advertencias=advertencias,
                    causas={d:[f['causa'] for f in filas if f['depto']==d and f['causa']!=TODAS]
                            for d in [nacional,*residencias]})
    return largo,estructura

def guardar_frecuencias(indicador, informe):
    # Contar filas originales, independientemente de la validación de defunciones.
    contador = Counter()
    for anio in ANIOS:
        _,_,matriz,_ = leer_hoja(anio,indicador[str(anio)])
        encabezado = unico([fila for fila in matriz if 'Departamento de residencia' in fila
                            and 'Causa de muerte' in fila and 'Total' in fila], 'Encabezado de frecuencia')
        cd,cc = encabezado.index('Departamento de residencia'),encabezado.index('Causa de muerte')
        apariciones_anio = 0
        for fila in matriz:
            d,c = fila[cd],fila[cc]
            if normalizar(d) not in ESPERADAS or c == TODAS:
                continue
            if not isinstance(c,str) or not c.strip():
                raise ValueError(f'{anio}: descripción vacía en bloque de residencia')
            contador[c] += 1
            apariciones_anio += 1
        informe.append(f'Frecuencia textual {anio}: {apariciones_anio} filas fuente; sin contar sexos ni agregados nacionales.')
    total_apariciones = sum(contador.values())
    esperadas = sorted(contador.items(), key=lambda p:(-p[1],p[0]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Frecuencia de causas'
    ws.append(['descripcion','veces_aparece'])
    for fila in esperadas:
        ws.append(fila)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions['A'].width = 100
    ws.column_dimensions['B'].width = 20
    for c in ws[1]:
        c.font = Font(bold=True)
    for fila in ws.iter_rows(min_row=2):
        fila[0].alignment = Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[fila[0].row].height = 32
    notas = wb.create_sheet('Criterio de conteo')
    for texto in [
        'Fuente: IND-004 en los 16 libros originales de 2009 a 2024; todas las residencias publicadas, incluidos Extranjero e Ignorado.',
        'veces_aparece = cantidad de filas fuente donde figura esa descripción en los bloques de residencia. Cada fila cuenta una vez; si la misma descripción aparece en dos filas con códigos distintos, ambas cuentan.',
        'No es la suma de defunciones, ni la cantidad de archivos distintos. Una descripción puede aparecer en varias residencias dentro del mismo libro.',
        'Se excluyen el bloque nacional y Todas las causas, para no duplicar agregados; Hombres y Mujeres no se cuentan como apariciones distintas.',
        'Se conserva cada descripción exacta, incluidas diferencias de escritura; no se fusionan enfermedades por similitud ni por código CIE.',
        'Incluye enfermedades, causas externas, síntomas y residuales como Otras causas; no es un listado exclusivo de enfermedades clínicas.',
        'Las listas son selectivas: una causa no listada no equivale a cero defunciones. El residual Otras causas puede cambiar de composición.',
        'Este reporte es de frecuencia textual, no de defunciones. Se permiten descripciones repetidas por instrucción del proyecto. La diferencia de una defunción del bloque nacional de 2023 sigue documentada, sin corregir.',
        f'Total de descripciones distintas: {len(esperadas)}; apariciones en filas fuente: {total_apariciones}.',
    ]:
        notas.append([texto])
    notas.column_dimensions['A'].width = 115
    for fila in notas:
        fila[0].alignment = Alignment(wrap_text=True,vertical='top')
        notas.row_dimensions[fila[0].row].height = 45
    FRECUENCIAS.parent.mkdir(parents=True,exist_ok=True)
    temporal = FRECUENCIAS.with_name(FRECUENCIAS.stem+'.tmp.xlsx')
    try:
        wb.save(temporal)
        comprobacion = openpyxl.load_workbook(temporal,read_only=True,data_only=True)
        try:
            assert list(comprobacion['Frecuencia de causas'].iter_rows(min_row=2,values_only=True)) == esperadas
        finally:
            comprobacion.close()
        temporal.replace(FRECUENCIAS)
    finally:
        wb.close()
        temporal.unlink(missing_ok=True)
    informe.append(f'Frecuencias verificadas: {len(esperadas)} descripciones; {total_apariciones} filas fuente. {FRECUENCIAS}')


def main():
    informe = ['VALIDACIÓN IND-004 — SERIE 2009–2024',
               'Formato: anio, depto_residencia, causa_muerte, sexo, conteo, tiene_cie10.',
               'Se conserva la clasificación publicada; no se inventan causas, códigos ni conteos.',
               'La validación de reproducción del cuadro no confirma equivalencia estadística de las causas.']
    resultados,estructuras = [],{}
    indicador = None
    frecuencia_generada = False
    try:
        with CATALOGO.open(encoding='utf-8-sig',newline='') as f:
            indicador = unico([r for r in csv.DictReader(f) if r['id_provisional']=='IND-004'],'Entrada IND-004')
        for anio in ANIOS:
            try:
                datos,estructura = transformar(anio,indicador[str(anio)],informe)
                if datos is not None:
                    resultados.append(datos);estructuras[anio]=estructura
            except Exception as error:
                informe.append(f'ERROR {anio}: {type(error).__name__}: {error}')
    except Exception as error:
        informe.append(f'ERROR de catálogo: {type(error).__name__}: {error}')
    if indicador is not None:
        try:
            guardar_frecuencias(indicador,informe)
            frecuencia_generada = True
        except Exception as error:
            informe.append(f'ERROR de frecuencias: {type(error).__name__}: {error}')
    informe.append(f'Años con reproducción verificada: {len(resultados)}/{len(ANIOS)}.')
    exito = len(resultados)==len(ANIOS)
    if exito:
        informe.append('\nDIFERENCIAS ENTRE AÑOS CONSECUTIVOS:')
        for anterior,actual in zip(ANIOS,ANIOS[1:]):
            for campo in ('dimensiones','encabezados','columnas','residencias'):
                informe.append(f'{anterior}→{actual}, {campo}: '+('iguales.' if estructuras[anterior][campo]==estructuras[actual][campo] else 'diferentes; revisar fuente.'))
            for d in estructuras[anterior]['residencias']:
                a = set(estructuras[anterior]['causas'][d])
                # Buscar la misma residencia con distinta escritura solo para el diagnóstico.
                equivalentes = [k for k in estructuras[actual]['residencias'] if normalizar(k)==normalizar(d)]
                if len(equivalentes)!=1:
                    informe.append(f'{anterior}→{actual}, {d}: residencia sin correspondencia única.')
                    continue
                b = set(estructuras[actual]['causas'][equivalentes[0]])
                informe.append(f'{anterior}→{actual}, {d}: solo listadas antes={sorted(a-b)!r}; después={sorted(b-a)!r}. No indican ausencia de defunciones.')
        datos = pd.concat(resultados,ignore_index=True).sort_values(['anio','depto_residencia','causa_muerte','sexo'],kind='stable').reset_index(drop=True)
        SALIDA.parent.mkdir(parents=True,exist_ok=True)
        temporal = SALIDA.with_suffix('.csv.tmp')
        try:
            datos.to_csv(temporal,index=False,encoding='utf-8-sig')
            pd.testing.assert_frame_equal(datos,pd.read_csv(temporal,encoding='utf-8-sig'))
            temporal.replace(SALIDA)
            estado = 'EXPORTADA CON ADVERTENCIAS EN LOS ORIGINALES' if any(e['advertencias'] for e in estructuras.values()) else 'REPRODUCCIÓN VERIFICADA'
            informe.append(f'\nRESULTADO GLOBAL: {estado}; {len(datos)} registros. {SALIDA}')
        except Exception as error:
            exito=False
            informe.append(f'ERROR exportando: {type(error).__name__}: {error}')
        finally:
            temporal.unlink(missing_ok=True)
    if not exito:
        informe.append('\nRESULTADO GLOBAL: NO VALIDADO; no se generó ni actualizó el CSV conjunto.')
        if SALIDA.exists():
            informe.append('Un CSV anterior no queda validado por esta ejecución.')
    REPORTE.parent.mkdir(parents=True,exist_ok=True)
    REPORTE.write_text('\n'.join(informe)+'\n',encoding='utf-8')
    for linea in informe:
        if linea.startswith(('\nAÑO:','Intersecciones','RESULTADO','\nRESULTADO','ERROR','Total:','Hombres:','Mujeres:')):
            print(linea)
    print(f'Reporte: {REPORTE}')
    if frecuencia_generada:
        print(f'Frecuencias: {FRECUENCIAS}')
    if not exito:
        raise SystemExit(1)


if __name__=='__main__':
    main()
