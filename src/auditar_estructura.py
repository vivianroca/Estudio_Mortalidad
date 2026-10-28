"""Auditoría estructural 2009–2024: candidatos para revisión, no equivalencias confirmadas."""

import csv
import re
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / 'reports' / 'estructura'
ANIOS = tuple(range(2009, 2025))
NO = 'No identificado'
DIMENSIONES = {
    'departamento residencia': r'departamento.{0,12}residencia',
    'departamento ocurrencia': r'departamento.{0,12}ocurrencia',
    'municipio': r'municipio', 'sexo': r'\bsexo\b|\bhombres\b|\bmujeres\b',
    'edad': r'\bedad\b|\bedades\b', 'edades simples': r'edades simples',
    'grupos de edad': r'grupos de edad', 'estado civil': r'estado civil',
    'causas de muerte': r'causas? de muerte', 'causas externas': r'causas? externas?',
    'certificación': r'certificacion', 'asistencia': r'asistencia',
    'lugar ocurrencia': r'lugar de ocurrencia', 'mes': r'\bmes\b',
    'día': r'\bdia\b', 'año ocurrencia': r'ano de ocurrencia',
    'área geográfica': r'area geografica', 'grupo étnico': r'grupo etnico',
}


def normalizar(texto):
    texto = ''.join(c for c in unicodedata.normalize('NFD', texto.casefold())
                    if unicodedata.category(c) != 'Mn')
    return ' '.join(texto.split())


def titulo_comparable(texto):
    # Solo quitar referencias temporales explícitas; conservar edades y categorías.
    texto = normalizar(texto)
    texto = re.sub(r'\b(?:ano|periodo)\s+\d{4}(?:\s*[-–]\s*\d{4})?', '', texto)
    return ' '.join(re.findall(r'\w+', texto))


def posicion(r, c):
    return f'{get_column_letter(c + 1)}{r + 1}'


def vista_xls(hoja):
    matriz = [[hoja.cell_value(r, c) for c in range(min(20, hoja.ncols))]
              for r in range(min(25, hoja.nrows))]
    combinadas = [f'{posicion(r1,c1)}:{posicion(r2-1,c2-1)}'
                  for r1,r2,c1,c2 in hoja.merged_cells if r1 < 25 and c1 < 20]
    return matriz, combinadas


def analizar(anio, archivo, nombre, filas, columnas, matriz, combinadas):
    observaciones = ['Ventana inspeccionada: filas 1–25, columnas A–T como máximo.',
                     'Dimensiones declaradas del libro; pueden incluir formato sin datos.']
    if combinadas:
        observaciones.append('Celdas combinadas (texto en ancla): ' + ', '.join(combinadas))
    textos = []
    for r, fila in enumerate(matriz):
        textos.append([(c, ' '.join(v.split())) for c,v in enumerate(fila)
                       if isinstance(v,str) and v.strip() and normalizar(v) not in
                       {'regresar','volver','inicio'} and not v.startswith('=')])
    indice = any(normalizar(t) in {'contenido','indice'} for fila in textos for _,t in fila)
    documental = normalizar(nombre) in {'nota tecnica','ficha tecnica','directorio','presentacion'}
    tipo_hoja = 'índice' if indice else ('documentación' if documental else 'cuadro candidato')
    titulo, subtitulo, encabezados, posiciones = NO, NO, [], []
    titulo_filas = set()
    if not indice and not documental:
        candidatos = [(r,c,t) for r,fila in enumerate(textos) for c,t in fila
                      if len(t) >= 40 and re.search(r'\bdefunciones\b',normalizar(t))
                      and len(fila) <= 2 and not any(
                          isinstance(v,(int,float)) for v in matriz[r])]
        if candidatos:
            r,c,t = candidatos[0]
            partes = [t]
            titulo_filas.add(r)
            # Continuaciones en filas siguientes con una sola celda descriptiva larga.
            for rr in range(r+1, min(r+4,len(textos))):
                if len(textos[rr]) == 1 and len(textos[rr][0][1]) >= 40:
                    partes.append(textos[rr][0][1]); titulo_filas.add(rr)
                else:
                    break
            titulo = ' '.join(partes)
            observaciones.append('Título candidato en ' + ', '.join(str(i+1) for i in sorted(titulo_filas)))
            fin_titulo = max(titulo_filas)
            # La primera fila numérica marca probable inicio de datos,
            # excepto secuencias explícitas de años bajo un título temporal.
            inicio_datos = len(matriz)
            for rr in range(fin_titulo+1,len(matriz)):
                numeros = [v for v in matriz[rr] if isinstance(v,(int,float))]
                anios = (numeros and 'año ocurrencia' in dimensiones(titulo)
                         and all(float(v).is_integer() and 1900 <= v <= 2100 for v in numeros))
                if numeros and not anios:
                    inicio_datos = rr
                    break
            for rr in range(fin_titulo+1,inicio_datos):
                if len(textos[rr]) == 1 and re.search(
                    r'^(guatemala|republica|ano |periodo |en |\()',normalizar(textos[rr][0][1])):
                    subtitulo = textos[rr][0][1] if subtitulo == NO else subtitulo+' | '+textos[rr][0][1]
                    continue
                for cc,t in textos[rr]:
                    encabezados.append(t); posiciones.append(f'{posicion(rr,cc)}: {t}')
                # Solo números que actúan como etiquetas de año; no registrar medidas.
                if 'año ocurrencia' in dimensiones(titulo):
                    for cc,v in enumerate(matriz[rr]):
                        if isinstance(v,(int,float)) and float(v).is_integer() and 1900 <= v <= 2100:
                            encabezados.append(str(int(v))); posiciones.append(f'{posicion(rr,cc)}: {int(v)}')
            if inicio_datos < len(matriz):
                observaciones.append(f'Posible inicio de datos: fila {inicio_datos+1}; solo se usaron tipos, no totales.')
    else:
        observaciones.append('Hoja de '+tipo_hoja+': excluida del catálogo de indicadores; conservada en inventario.')
    if titulo == NO:
        observaciones.append('Título sin confianza suficiente: revisión manual.')
    if not encabezados:
        observaciones.append('Encabezados no identificados con confianza.')
    fuente = ' '.join([titulo if titulo != NO else '', *encabezados])
    unidades = [t for fila in textos for _,t in fila if re.search(
        r'porcentaje|por ciento|\btasa\b|por (?:cada )?(?:mil|cien)|unidad(?:es)?|numero de defunciones',normalizar(t))]
    observaciones.append('Unidad explícita candidata: ' + (' | '.join(dict.fromkeys(unidades)) or NO))
    dims = dimensiones(fuente)
    return dict(anio=anio, archivo=archivo, hoja=nombre, tipo_hoja=tipo_hoja, filas=filas, columnas=columnas,
                titulo=titulo, subtitulo=subtitulo, encabezados=' | '.join(encabezados) or NO,
                posicion_encabezados=' | '.join(posiciones) or NO,
                dimensiones=' | '.join(sorted(dims)) or NO,
                observaciones=' '.join(observaciones))


def dimensiones(texto):
    texto = normalizar(texto)
    return {nombre for nombre,patron in DIMENSIONES.items() if re.search(patron,texto)}


def inspeccionar(anio, errores):
    rutas = [RAIZ/'data'/'raw_data'/f'{anio}{ext}' for ext in ('.xls','.xlsx')
             if (RAIZ/'data'/'raw_data'/f'{anio}{ext}').is_file()]
    if len(rutas) != 1:
        errores.append(f'{anio}: se esperaba un archivo; encontrados {len(rutas)}')
        return [], False
    ruta = rutas[0]
    resultados = []
    libro = None
    try:
        if ruta.suffix == '.xls':
            libro = xlrd.open_workbook(str(ruta), on_demand=True, formatting_info=True)
            nombres = libro.sheet_names()
        else:
            libro = openpyxl.load_workbook(ruta,read_only=True,data_only=False,keep_links=False)
            nombres = libro.sheetnames
        for nombre in nombres:
            try:
                if ruta.suffix == '.xls':
                    hoja = libro.sheet_by_name(nombre)
                    matriz, combinadas = vista_xls(hoja)
                    filas,columnas = hoja.nrows,hoja.ncols
                else:
                    hoja = libro[nombre]
                    filas,columnas = hoja.max_row or 0,hoja.max_column or 0
                    matriz = list(hoja.iter_rows(min_row=1,max_row=min(25,filas),
                                               max_col=min(20,columnas),values_only=True)) if filas and columnas else []
                    # ReadOnlyWorksheet no expone rangos combinados: leer solo metadatos XML.
                    combinadas = []
                    with ZipFile(ruta) as z:
                        with z.open(hoja._worksheet_path.lstrip('/')) as xml:
                            for _,el in ElementTree.iterparse(xml,events=('end',)):
                                if el.tag.endswith('}mergeCell'):
                                    ref = el.attrib['ref']
                                    from openpyxl.utils.cell import range_boundaries
                                    c1,r1,_,_ = range_boundaries(ref)
                                    if r1 <= 25 and c1 <= 20:
                                        combinadas.append(ref)
                                el.clear()
                resultados.append(analizar(anio,ruta.name,nombre,filas,columnas,matriz,combinadas))
            except Exception as error:
                errores.append(f'{anio} / {nombre}: {type(error).__name__}: {error}')
        return resultados, True
    except Exception as error:
        errores.append(f'{anio}: {type(error).__name__}: {error}')
        return resultados, False
    finally:
        if libro is not None:
            libro.release_resources() if ruta.suffix == '.xls' else libro.close()


def jaccard(a,b):
    return len(a & b)/len(a | b) if a and b else 0.0


def comparar(a,b):
    ta,tb = titulo_comparable(a['titulo']),titulo_comparable(b['titulo'])
    st = SequenceMatcher(None,ta,tb).ratio() if NO not in (a['titulo'],b['titulo']) else 0
    ha = {normalizar(t) for t in a['encabezados'].split(' | ')} - {normalizar(NO)}
    hb = {normalizar(t) for t in b['encabezados'].split(' | ')} - {normalizar(NO)}
    sh = jaccard(ha,hb)
    da,db = dimensiones(a['titulo']),dimensiones(b['titulo'])
    sd = jaccard(da,db)
    forma = min(a['columnas'],b['columnas'])/max(a['columnas'],b['columnas'],1)
    motivos = [f'Títulos={st:.3f}; encabezados={sh:.3f}; dimensiones={sd:.3f}; razón columnas={forma:.3f}.']
    # Evitar que vocabulario compartido oculte cambios de población/granularidad.
    def poblacion(titulo):
        texto = normalizar(titulo)
        if re.search(r'post[ -]?neonatal', texto):
            return 'postneonatal'
        if 'neonatal' in texto:
            return 'neonatal'
        return 'infantil' if 'infantiles' in texto else 'general'

    conflicto = poblacion(a['titulo']) != poblacion(b['titulo']) or da != db
    if st >= .93 and sh >= .75 and da == db and da and forma >= .9 and not conflicto:
        clase = 'Posible equivalencia alta'
    elif st >= .8 and sh >= .4 and sd >= .6 and not conflicto:
        clase = 'Posible equivalencia media'
    else:
        clase = 'Equivalencia no determinada'
    if conflicto:
        motivos.append('Diferencias en población o granularidad; revisión manual.')
    if (a['filas'],a['columnas']) != (b['filas'],b['columnas']):
        motivos.append(f"Dimensiones cambiaron: {a['filas']}x{a['columnas']} → {b['filas']}x{b['columnas']}.")
    if ta == tb and a['titulo'] != b['titulo']:
        motivos.append('Coinciden al normalizar escritura y referencias temporales explícitas.')
    return dict(anio_origen=a['anio'],hoja_origen=a['hoja'],titulo_origen=a['titulo'],
                anio_destino=b['anio'],hoja_destino=b['hoja'],titulo_destino=b['titulo'],
                similitud_titulos=round(st,4),similitud_encabezados=round(sh,4),
                clasificacion=clase,motivo=' '.join(motivos))


def guardar_csv(nombre,filas,campos):
    with (SALIDA/nombre).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=campos)
        writer.writeheader(); writer.writerows(filas)


def catalogar(inventario, pares):
    """Componentes candidatos; solo agrupar si toda pareja tiene respaldo directo.

    Un componente incompatible se conserva como fichas separadas ambiguas,
    sin elegir una rama ni propagar transitivamente una similitud débil.
    """
    por_clave = {(f['anio'], f['hoja']): f for f in inventario if f['titulo'] != NO}
    vecinos = {clave: set() for clave in por_clave}
    evidencia = {}
    for p in pares:
        if p['clasificacion'] == 'Equivalencia no determinada':
            continue
        a = (p['anio_origen'], p['hoja_origen'])
        b = (p['anio_destino'], p['hoja_destino'])
        vecinos[a].add(b); vecinos[b].add(a)
        evidencia[frozenset((a, b))] = p
    componentes, pendientes = [], set(por_clave)
    while pendientes:
        inicio = min(pendientes)
        componente, cola = set(), [inicio]
        while cola:
            clave = cola.pop()
            if clave in componente:
                continue
            componente.add(clave)
            cola.extend(vecinos[clave] - componente)
        pendientes -= componente
        componentes.append(sorted(componente))

    grupos = []
    for componente in componentes:
        motivos = []
        anios = Counter(a for a, _ in componente)
        if any(n > 1 for n in anios.values()):
            motivos.append('Dos o más hojas del mismo año compiten por un indicador; no se selecciona ninguna rama.')
        # No basta A~B~C: también exigir evidencia directa A~C, etc.
        for i, a in enumerate(componente):
            for b in componente[i+1:]:
                if a[0] == b[0]:
                    continue
                p = evidencia.get(frozenset((a,b))) or comparar(por_clave[a],por_clave[b])
                if p['clasificacion'] == 'Equivalencia no determinada':
                    motivos.append(f'Cadena incompatible: {a[0]}/{a[1]} frente a {b[0]}/{b[1]}.')
        if motivos:
            # Conservar toda la evidencia, pero no formar un grupo sin respaldo.
            for clave in componente:
                grupos.append(([clave], 'ambiguo', motivos, componente))
        else:
            estado = 'pendiente' if len(componente) > 1 else 'sin candidato'
            grupos.append((componente, estado, [], componente))

    catalogo = []
    for numero, (grupo, estado, motivos, relacionados) in enumerate(grupos, 1):
        primero = por_clave[grupo[0]]  # Referencia cronológica, no elección entre competidores.
        observaciones = ['Agrupación provisional; equivalencia estadística no confirmada.']
        if len(grupo) > 1:
            observaciones.append('Conectada por años consecutivos y verificada entre todos los miembros; no se usó el nombre de hoja.')
        for i, clave in enumerate(grupo):
            if i:
                anterior = por_clave[grupo[i-1]]
                actual = por_clave[clave]
                p = comparar(anterior,actual)
                if (anterior['filas'],anterior['columnas']) != (actual['filas'],actual['columnas']):
                    observaciones.append(f'{anterior["anio"]}→{actual["anio"]}: '+p['motivo']+
                                         ' Cambio de presentación posible; no implica cambio de definición.')
                if p['clasificacion'] == 'Posible equivalencia media':
                    observaciones.append(f'{anterior["anio"]}→{actual["anio"]}: evidencia media, revisar encabezados y definición.')
        if estado == 'sin candidato':
            observaciones.append('Sin coincidencia consecutiva suficientemente respaldada; las demás columnas no indican ausencia del indicador.')
        if motivos:
            observaciones.extend(motivos)
            observaciones.append('Relacionados sin agrupar: '+ '; '.join(f'{a}/{h}' for a,h in relacionados))
        fila = dict(id_provisional=f'IND-{numero:03d}',
                    titulo_representativo=primero['titulo'],
                    dimensiones=' | '.join(sorted(set().union(*[
                        dimensiones(por_clave[k]['titulo']+' '+por_clave[k]['encabezados']) for k in grupo]))) or NO)
        for anio in ANIOS:
            nombres = [h for a,h in grupo if a == anio]
            fila[str(anio)] = nombres[0] if len(nombres) == 1 else NO
        fila.update(cantidad_anios=len({a for a,_ in grupo}), estado_revision=estado,
                    observaciones=' '.join(observaciones))
        catalogo.append(fila)
    return catalogo


def main():
    errores, inventario = [], []
    archivos = 0
    for anio in ANIOS:
        filas,ok = inspeccionar(anio,errores)
        inventario.extend(filas); archivos += int(ok)
        print(f'{anio}: {len(filas)} hojas examinadas; archivo abierto: {ok}', flush=True)
    pares = [comparar(a,b) for a in inventario for b in inventario
             if b['anio'] == a['anio']+1]
    pares.sort(key=lambda p:(-p['similitud_titulos'],-p['similitud_encabezados']))
    SALIDA.mkdir(parents=True,exist_ok=True)
    guardar_csv('inventario_estructura.csv',inventario,
                ['anio','archivo','hoja','tipo_hoja','filas','columnas','titulo','subtitulo','encabezados',
                 'posicion_encabezados','dimensiones','observaciones'])
    guardar_csv('posibles_equivalencias.csv',pares,
                ['anio_origen','hoja_origen','titulo_origen','anio_destino','hoja_destino','titulo_destino',
                 'similitud_titulos','similitud_encabezados','clasificacion','motivo'])
    catalogo = catalogar(inventario,pares)
    guardar_csv('catalogo_indicadores.csv',catalogo,
                ['id_provisional','titulo_representativo','dimensiones',*map(str,ANIOS),
                 'cantidad_anios','estado_revision','observaciones'])
    identificados = sum(f['titulo'] != NO for f in inventario)
    completos = [f for f in catalogo if f['cantidad_anios'] == len(ANIOS)]
    parciales = [f for f in catalogo if 1 < f['cantidad_anios'] < len(ANIOS)]
    ambiguos = [f for f in catalogo if f['estado_revision'] == 'ambiguo']
    solitarios = [f for f in catalogo if f['estado_revision'] == 'sin candidato']
    lineas = ['AUDITORÍA ESTRUCTURAL: 2009–2024',
              f'Archivos abiertos: {archivos}/16; hojas examinadas: {len(inventario)}.',
              f'Hojas con título identificado: {identificados}; sin título: {len(inventario)-identificados}.',
              f'Fichas de indicadores provisionales: {len(catalogo)}.',
              f'Grupos longitudinales candidatos (2 o más años): {len(completos)+len(parciales)}.',
              f'Candidatos con cobertura en los 16 años: {len(completos)}.',
              f'Candidatos con cobertura parcial (2–15 años): {len(parciales)}.',
              f'Fichas ambiguas sin agrupar: {len(ambiguos)}; fichas sin candidato: {len(solitarios)}.',
              f'Errores registrados: {len(errores)}.',
              '\nMÉTODO Y LIMITACIONES',
              'Ventana de 25 filas y 20 columnas; títulos, encabezados, posiciones, dimensiones y metadatos de celdas combinadas.',
              'Comparaciones de evidencia principal solo entre años consecutivos. No se usan nombres de hojas para puntuar.',
              'Similitud títulos: SequenceMatcher; encabezados y dimensiones: Jaccard (0–1).',
              'Alta: título ≥0.93, encabezados ≥0.75, mismas dimensiones detectadas, razón columnas ≥0.9, sin conflicto poblacional.',
              'Media: título ≥0.8, encabezados ≥0.4, mismas dimensiones detectadas, sin conflicto poblacional.',
              'Agrupar exige compatibilidad directa de todas las parejas, incluso no consecutivas como control antitransitividad.',
              'Si un componente tiene competidores del mismo año o parejas incompatibles, se divide en fichas individuales ambiguas; no se elige rama.',
              'Las fichas ambiguas incluyen todas las hojas relacionadas en observaciones. Ningún vínculo confirma equivalencia estadística.',
              'Las hojas sin título permanecen en inventario y resumen, sin inventar indicadores.',
              'Índices, presentación, directorio y notas/fichas técnicas se conservan como documentación, sin catalogarlos como indicadores.',
              'No identificado significa evidencia insuficiente; nunca ausencia confirmada.',
              'Dimensiones declaradas pueden incluir formato; cambios de tamaño no prueban cambios de definición.',
              'No se comparan cifras de mortalidad. La detección es heurística y puede omitir títulos/encabezados.',
              '\nCLASIFICACIONES DE PARES: '+str(dict(Counter(p['clasificacion'] for p in pares))),
              '\nCATÁLOGO / COBERTURA:']
    for f in catalogo:
        lineas.append(f"- {f['id_provisional']}: {f['cantidad_anios']} años; {f['estado_revision']}; {f['titulo_representativo']}")
    lineas.append('\nCAMBIOS ESTRUCTURALES Y CASOS AMBIGUOS:')
    for f in catalogo:
        if f['estado_revision'] == 'ambiguo' or 'Cambio de presentación' in f['observaciones']:
            lineas.append(f"- {f['id_provisional']}: {f['observaciones']}")
    lineas.append('\nHOJAS SIN TÍTULO IDENTIFICADO:')
    lineas.extend(f"- {f['anio']}/{f['hoja']}: {f['observaciones']}" for f in inventario if f['titulo'] == NO)
    lineas.append('\nERRORES:')
    lineas.extend(errores or ['Ninguno.'])
    (SALIDA/'resumen_estructura.txt').write_text('\n'.join(lineas)+'\n',encoding='utf-8')
    print('\n'.join(lineas[:9]))
    print(f'Reportes: {SALIDA}')


if __name__ == '__main__':
    main()
