"""IND-006 de 2009–2024: residencia, grupo de edad original y sexo en formato largo."""

import csv
import hashlib
from pathlib import Path

import pandas as pd
import xlrd
import openpyxl

from transformar_ind002 import DEPARTAMENTOS, entero, normalizar, unico

RAIZ = Path(__file__).resolve().parents[1]
CATALOGO = RAIZ / 'reports/estructura/catalogo_indicadores.csv'
SALIDA = RAIZ / 'data/processed/ind006_2009_2024.csv'
REPORTE = RAIZ / 'reports/estructura/validacion_ind006_2009_2024.txt'
ANIOS = tuple(range(2009, 2025))
COLUMNAS = ['anio', 'depto_residencia', 'grupo_edad', 'sexo', 'conteo']
RESIDENCIAS = (*DEPARTAMENTOS, 'Extranjero', 'Ignorado')
EDADES = ('Menos de 1 año', '  1 - 4  ', '  5 - 9  ', *(f'{n} - {n+4}' for n in range(10,95,5)), '95 y más', 'Ignorado')
SEXOS = ('Hombres', 'Mujeres')
NACIONAL = 'Todos los departamentos'
TODAS = 'Todas las edades'


def transformar(anio, nombre, informe):
    rutas = [RAIZ / 'data/raw_data' / f'{anio}{ext}' for ext in ('.xls', '.xlsx')
             if (RAIZ / 'data/raw_data' / f'{anio}{ext}').is_file()]
    ruta = unico(rutas, f'Archivo del año {anio}')
    huella = hashlib.sha256(ruta.read_bytes()).hexdigest()
    if ruta.suffix == '.xls':
        libro = xlrd.open_workbook(str(ruta), on_demand=True)
        try:
            hoja = libro.sheet_by_name(nombre)
            matriz = [hoja.row_values(r) for r in range(hoja.nrows)]
            dimensiones = (hoja.nrows, hoja.ncols)
        finally:
            libro.release_resources()
    else:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=False, keep_links=False)
        try:
            hoja = libro[nombre]
            dimensiones = (hoja.max_row, hoja.max_column)
            matriz = [[v if v is not None else '' for v in fila]
                      for fila in hoja.iter_rows(values_only=True)]
        finally:
            libro.close()
    informe.extend([f'\nAÑO: {anio}', f'Archivo: {ruta.name}; hoja del catálogo: {nombre}',
                    f'Dimensiones: {dimensiones[0]} filas × {dimensiones[1]} columnas.',
                    f'SHA-256: {huella}'])
    errores = []
    r_enc = unico([r for r, fila in enumerate(matriz)
                   if all(n in fila for n in ('Departamento de residencia', 'Grupos de edad', 'Total', *SEXOS))],
                  'Fila de encabezados')
    encabezados = matriz[r_enc]
    cols = {n: unico([c for c, v in enumerate(encabezados) if v == n], f'Columna {n}')
            for n in ('Departamento de residencia', 'Grupos de edad', 'Total', *SEXOS)}
    informe.append(f'Encabezados en fila {r_enc+1}: {cols} (columnas indexadas desde cero).')
    nacional = unico(list(dict.fromkeys(
        fila[cols['Departamento de residencia']] for fila in matriz[r_enc+1:]
        if normalizar(fila[cols['Departamento de residencia']]) in
        {'todos los departamentos', 'total pais'})), 'Etiqueta del agregado nacional')
    informe.append(f'Agregado nacional identificado: {nacional!r}; se usa solo para validación.')
    valores, origen = {}, {}
    notas = []
    for r in range(r_enc+1, len(matriz)):
        fila = matriz[r]
        if not any(v != '' for v in fila):
            continue
        depto = fila[cols['Departamento de residencia']]
        edad = fila[cols['Grupos de edad']]
        if isinstance(depto, str) and normalizar(depto).startswith(('fuente:', 'nota:', 'notas:')):
            notas.append(f'Fila {r+1}: {depto}')
            if any(v != '' for c,v in enumerate(fila) if c != cols['Departamento de residencia']):
                errores.append(f'Fila {r+1}: nota con celdas adicionales; revisión necesaria')
            continue
        if depto not in (nacional, *RESIDENCIAS) or edad not in (TODAS, *EDADES):
            errores.append(f'Fila {r+1}: categoría desconocida o posible subtotal: {depto!r}, {edad!r}')
            continue
        clave = (depto, edad)
        if clave in valores:
            errores.append(f'Fila {r+1}: clave duplicada {clave!r}')
            continue
        if any(v != '' for c,v in enumerate(fila) if c not in cols.values()):
            errores.append(f'Fila {r+1}: celdas fuera de las columnas identificadas')
        datos = {}
        for sexo in ('Total', *SEXOS):
            c = cols[sexo]
            try:
                datos[sexo] = entero(fila[c], f'fila {r+1}, columna {c+1}')
            except ValueError as error:
                errores.append(str(error))
        valores[clave] = datos
        origen[clave] = r+1
    esperadas = {(d,e) for d in (nacional,*RESIDENCIAS) for e in (TODAS,*EDADES)}
    faltantes = esperadas - set(valores)
    if faltantes:
        errores.append(f'Combinaciones residencia/edad faltantes: {sorted(faltantes)!r}')
    informe.extend([f'Categorías: {len(RESIDENCIAS)} residencias, {len(EDADES)} edades y {len(SEXOS)} sexos.',
                    'Residencias: '+', '.join(RESIDENCIAS),
                    'Edades originales: '+', '.join(EDADES),
                    f'Filas nacionales: {[origen[k] for k in origen if k[0]==nacional]}.',
                    f'Filas de todas las edades: {[origen[k] for k in origen if k[1]==TODAS]}.',
                    'Se excluyen ambos tipos de filas y la columna Total del CSV, conservándolos para validación.',
                    'Extranjero e Ignorado son categorías originales, no subtotales; se conservan.',
                    'Se conservan todas las etiquetas de grupos de edad tal como aparecen en el Excel, sin recodificarlas.',
                    'Notas/fuentes: '+('; '.join(notas) or 'ninguna dentro de la hoja.'),
                    f'Claves originales distintas: {len(valores)}; esperadas: {len(esperadas)}.',
                    'No se imputan guiones, blancos ni no numéricos; cualquier caso bloquea la salida.'])
    if errores:
        informe.extend('ERROR: '+e for e in errores)
        return None, None

    ancho = pd.DataFrame([{'depto_residencia':d, 'grupo_edad':e, **{s:valores[d,e][s] for s in SEXOS}}
                          for d in RESIDENCIAS for e in EDADES])
    largo = ancho.melt(id_vars=['depto_residencia','grupo_edad'], value_vars=list(SEXOS),
                       var_name='sexo',value_name='conteo')
    largo.insert(0,'anio',anio)
    largo = largo[COLUMNAS]
    largo['conteo'] = largo['conteo'].astype('int64')
    original = {(d,e,s):valores[d,e][s] for d in RESIDENCIAS for e in EDADES for s in SEXOS}
    transformado = {(v.depto_residencia,v.grupo_edad,v.sexo):v.conteo for v in largo.itertuples(index=False)}
    if transformado != original or largo.duplicated(COLUMNAS[:-1]).any() or len(largo)!=len(original):
        errores.append('Intersecciones incorrectas, faltantes o duplicadas tras melt')
    # Comprobar todas las filas, incluidos los márgenes, no solo las analíticas.
    fallos_fila = 0
    for (d,e), v in valores.items():
        if v['Hombres']+v['Mujeres'] != v['Total']:
            fallos_fila += 1
            errores.append(f'Fila {origen[d,e]}: suma de sexos no coincide con Total')
    informe.append(f'Validación Hombres+Mujeres=Total: {len(valores)} filas; diferencias: {fallos_fila}.')
    agrupados = largo.groupby(['depto_residencia','grupo_edad'])['conteo'].sum()
    for d in RESIDENCIAS:
        for e in EDADES:
            if int(agrupados[d,e]) != valores[d,e]['Total']:
                errores.append(f'Total por residencia y edad no coincide: {d}/{e}')
    informe.append('\nTOTALES POR RESIDENCIA Y SEXO (sumar edades, excluyendo Todas las edades):')
    por_depto = largo.groupby(['depto_residencia','sexo'])['conteo'].sum()
    for d in RESIDENCIAS:
        for sexo in ('Total', *SEXOS):
            calculado = sum(int(por_depto[d,s]) for s in SEXOS) if sexo=='Total' else int(por_depto[d,sexo])
            publicado = valores[d,TODAS][sexo]
            informe.append(f'{d} / {sexo}: publicado={publicado}; transformado={calculado}; diferencia={calculado-publicado}.')
            if calculado != publicado:
                errores.append(f'Total por residencia/sexo no coincide: {d}/{sexo}')
    informe.append('\nTOTALES nacionalES POR EDAD Y SEXO (sumar residencias):')
    por_edad = largo.groupby(['grupo_edad','sexo'])['conteo'].sum()
    for e in EDADES:
        for sexo in ('Total',*SEXOS):
            calculado = sum(int(por_edad[e,s]) for s in SEXOS) if sexo=='Total' else int(por_edad[e,sexo])
            publicado = valores[nacional,e][sexo]
            informe.append(f'{e} / {sexo}: publicado={publicado}; transformado={calculado}; diferencia={calculado-publicado}.')
            if calculado != publicado:
                errores.append(f'Total por edad/sexo no coincide: {e}/{sexo}')
    informe.append('\nTOTALES GENERALES:')
    for sexo in ('Total',*SEXOS):
        calculado = int(largo['conteo'].sum() if sexo=='Total' else largo.loc[largo.sexo==sexo,'conteo'].sum())
        publicado = valores[nacional,TODAS][sexo]
        informe.append(f'{sexo}: publicado={publicado}; transformado={calculado}; diferencia={calculado-publicado}.')
        if calculado != publicado:
            errores.append(f'Total general no coincide: {sexo}')
    if hashlib.sha256(ruta.read_bytes()).hexdigest()!=huella:
        errores.append('El archivo original cambió durante la ejecución')
    informe.extend([f'Intersecciones verificadas: {len(original)}; registros: {len(largo)}.',
                    'Guiones, vacíos o valores no numéricos en datos y márgenes: 0.',
                    'EJEMPLOS:',largo.head(6).to_csv(index=False).strip()])
    if errores:
        informe.extend('ERROR: '+e for e in errores)
        return None,None
    informe.append(f'RESULTADO {anio}: VALIDADO.')
    estructura = dict(dimensiones=dimensiones, encabezados=encabezados, columnas=cols,
                      categorias=list(valores), posiciones=origen)
    return largo,estructura


def main():
    informe = ['VALIDACIÓN IND-006 — SERIE 2009–2024',
               'Formato largo: anio, depto_residencia, grupo_edad, sexo, conteo.',
               'Solo se abre IND-006 en los años 2009–2024. Residencia y sexo conservan sus nombres originales.',
               'Grupos de edad: conservar exactamente las etiquetas originales, incluidos espacios, Menos de 1 año, 95 y más e Ignorado; no convertir intervalos a edades puntuales.']
    resultados, estructuras = [],{}
    try:
        with CATALOGO.open(encoding='utf-8-sig',newline='') as f:
            indicador = unico([r for r in csv.DictReader(f) if r['id_provisional']=='IND-006'], 'Entrada IND-006')
        for anio in ANIOS:
            try:
                datos,estructura = transformar(anio,indicador[str(anio)],informe)
                if datos is not None:
                    resultados.append(datos); estructuras[anio]=estructura
            except Exception as error:
                informe.append(f'ERROR {anio}: {type(error).__name__}: {error}')
    except Exception as error:
        informe.append(f'ERROR de catálogo: {type(error).__name__}: {error}')
    exito = len(resultados)==len(ANIOS)
    if exito:
        informe.append('\nDIFERENCIAS DE ESTRUCTURA ENTRE AÑOS CONSECUTIVOS:')
        for anterior,actual in zip(ANIOS,ANIOS[1:]):
            diferencias=[k for k in estructuras[anterior] if estructuras[anterior][k]!=estructuras[actual][k]]
            informe.append(f'{anterior}→{actual}: '+(', '.join(diferencias) if diferencias else
                           'sin diferencias en dimensiones, encabezados, posiciones, categorías o nombres.'))
        datos = pd.concat(resultados,ignore_index=True).sort_values(COLUMNAS[:-1],kind='stable').reset_index(drop=True)
        SALIDA.parent.mkdir(parents=True,exist_ok=True)
        temporal=SALIDA.with_suffix('.csv.tmp')
        try:
            datos.to_csv(temporal,index=False,encoding='utf-8-sig')
            pd.testing.assert_frame_equal(datos,pd.read_csv(temporal,encoding='utf-8-sig'))
            temporal.replace(SALIDA)
            informe.append(f'\nRESULTADO GLOBAL: SERIE VALIDADA; {len(datos)} registros. CSV releído y verificado: {SALIDA}')
        except Exception as error:
            exito=False
            informe.append(f'ERROR de exportación: {type(error).__name__}: {error}')
        finally:
            temporal.unlink(missing_ok=True)
    if not exito:
        informe.append('\nRESULTADO GLOBAL: NO VALIDADO. No se generó ni actualizó el CSV conjunto.')
        if SALIDA.exists():
            informe.append('El CSV existente corresponde a otra ejecución y no está validado por esta prueba.')
    REPORTE.parent.mkdir(parents=True,exist_ok=True)
    REPORTE.write_text('\n'.join(informe)+'\n',encoding='utf-8')
    for linea in informe:
        if linea.startswith(('\nAÑO:','RESULTADO','\nRESULTADO','ERROR','Intersecciones','Total:','Hombres:','Mujeres:')):
            print(linea)
    print(f'Reporte: {REPORTE}')
    if not exito:
        raise SystemExit(1)


if __name__=='__main__':
    main()
