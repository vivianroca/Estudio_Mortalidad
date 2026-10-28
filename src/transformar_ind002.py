"""Transformación reproducible de IND-002 en los libros de 2009 a 2024."""

import csv
import hashlib
import math
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
import xlrd

RAIZ = Path(__file__).resolve().parents[1]
CATALOGO = RAIZ / 'reports' / 'estructura' / 'catalogo_indicadores.csv'
SALIDA = RAIZ / 'data' / 'processed' / 'ind002_2009_2024.csv'
REPORTE = RAIZ / 'reports' / 'estructura' / 'validacion_ind002_2009_2024.txt'
ANIOS = tuple(range(2009, 2025))
COLUMNAS = ['anio', 'depto_ocurrencia', 'depto_residencia', 'conteo']
DEPARTAMENTOS = (
    'Guatemala', 'El Progreso', 'Sacatepéquez', 'Chimaltenango', 'Escuintla',
    'Santa Rosa', 'Sololá', 'Totonicapán', 'Quetzaltenango', 'Suchitepéquez',
    'Retalhuleu', 'San Marcos', 'Huehuetenango', 'Quiché', 'Baja Verapaz',
    'Alta Verapaz', 'Petén', 'Izabal', 'Zacapa', 'Chiquimula', 'Jalapa', 'Jutiapa',
)
ESPECIALES = ('Extranjero', 'Ignorado')


def normalizar(texto):
    texto = unicodedata.normalize('NFD', str(texto).casefold())
    return ' '.join(''.join(c for c in texto if not unicodedata.combining(c)).split())


def celda(r, c):
    return f'{get_column_letter(c+1)}{r+1}'


def unico(elementos, descripcion):
    if len(elementos) != 1:
        raise ValueError(f'{descripcion}: se esperaba uno; encontrados {elementos!r}')
    return elementos[0]


def entero(valor, posicion):
    if (isinstance(valor, bool) or not isinstance(valor, (int, float))
            or not math.isfinite(valor) or valor < 0 or int(valor) != valor):
        tipo = ('vacío' if valor in ('', None) else
                'guion pendiente de interpretación' if str(valor).strip() in ('-', '–', '—')
                else 'no numérico, negativo o no entero')
        raise ValueError(f'{posicion}: {tipo}: {valor!r}; no se imputó ningún valor')
    return int(valor)


def validar_nombres(nombres, esperados, descripcion, lineas, errores):
    duplicados = [n for n, cantidad in Counter(nombres).items() if cantidad > 1]
    claves = [normalizar(n) for n in nombres]
    variantes = [n for n, cantidad in Counter(claves).items() if cantidad > 1]
    faltantes = sorted(set(esperados) - set(nombres))
    adicionales = sorted(set(nombres) - set(esperados))
    lineas.append(f'{descripcion}: {len(nombres)} categorías; faltantes={faltantes}; '
                  f'adicionales={adicionales}; duplicados={duplicados}; '
                  f'colisiones al normalizar={variantes}.')
    if duplicados or variantes or faltantes or adicionales:
        errores.append(f'{descripcion}: categorías faltantes, duplicadas o inconsistentes')


def transformar_anio(anio, hoja_nombre, lineas):
    rutas = [RAIZ / 'data' / 'raw_data' / f'{anio}{ext}' for ext in ('.xls', '.xlsx')
             if (RAIZ / 'data' / 'raw_data' / f'{anio}{ext}').is_file()]
    ruta = unico(rutas, f'Archivo del año {anio}')
    hash_original = hashlib.sha256(ruta.read_bytes()).hexdigest()
    if ruta.suffix == '.xls':
        libro = xlrd.open_workbook(str(ruta), on_demand=True)
        try:
            hoja = libro.sheet_by_name(hoja_nombre)
            matriz = [hoja.row_values(r) for r in range(hoja.nrows)]
            dimensiones = (hoja.nrows, hoja.ncols)
        finally:
            libro.release_resources()
    else:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=False, keep_links=False)
        try:
            hoja = libro[hoja_nombre]
            dimensiones = (hoja.max_row, hoja.max_column)
            matriz = [[v if v is not None else '' for v in fila]
                      for fila in hoja.iter_rows(values_only=True)]
        finally:
            libro.close()
    lineas.extend([f'\nAÑO: {anio}', f'Archivo: {ruta.name}; hoja: {hoja_nombre}',
                   f'Dimensiones originales: {dimensiones[0]} filas × {dimensiones[1]} columnas.',
                   f'SHA-256 original: {hash_original}'])
    errores = []
    # Localizar cada estructura por etiquetas del propio libro, no por posiciones fijas.
    r_ejes, c_res = unico([(r,c) for r,fila in enumerate(matriz)
                          for c,v in enumerate(fila)
                          if normalizar(v) == 'departamento de residencia'], 'Eje residencia')
    ejes_oc = [(r,c) for r,fila in enumerate(matriz) for c,v in enumerate(fila)
               if normalizar(v) == 'departamento de ocurrencia']
    r_oc, c_oc = unico(ejes_oc, 'Eje ocurrencia')
    if r_oc != r_ejes:
        raise ValueError('Los ejes no están en la misma fila; revisar la estructura')
    r_enc = unico([r for r in range(r_ejes+1,len(matriz))
                   if any(normalizar(v) == 'total' for v in matriz[r])
                   and 'Guatemala' in matriz[r]], 'Fila de encabezados de ocurrencia')
    c_total = unico([c for c,v in enumerate(matriz[r_enc]) if normalizar(v) == 'total'],
                    'Columna de totales')
    cols = [c for c,v in enumerate(matriz[r_enc]) if isinstance(v,str) and v.strip()
            and c != c_total and c != c_res]
    ocurrencias = [matriz[r_enc][c] for c in cols]
    r_total = unico([r for r,fila in enumerate(matriz)
                     if normalizar(fila[c_res]) in {'todos los departamentos', 'total pais'}], 'Fila total general')
    if not r_ejes < r_enc < r_total or c_oc != c_total:
        raise ValueError('Orden de encabezados/totales o asociación del eje ocurrencia inesperados')
    filas = []
    notas = []
    for r in range(r_total+1,len(matriz)):
        nombre = matriz[r][c_res]
        if not any(v != '' for v in matriz[r]):
            continue
        if nombre in (*DEPARTAMENTOS, *ESPECIALES):
            filas.append(r)
        elif isinstance(nombre,str) and normalizar(nombre).startswith(('fuente:', 'nota:', 'notas:')):
            notas.append(f'Fila {r+1}: {nombre}')
        else:
            errores.append(f'Fila no clasificada / posible subtotal en {r+1}: {matriz[r]!r}')
    residencias = [matriz[r][c_res] for r in filas]
    validar_nombres(ocurrencias, DEPARTAMENTOS, 'Ocurrencia', lineas, errores)
    validar_nombres(residencias, (*DEPARTAMENTOS,*ESPECIALES), 'Residencia', lineas, errores)
    lineas.extend([f'Fila de ejes: {r_ejes+1}; fila de encabezados: {r_enc+1}.',
                   'Departamentos de ocurrencia: '+', '.join(f'{celda(r_enc,c)}={matriz[r_enc][c]}' for c in cols),
                   f'Totales por residencia: columna {chr(65+c_total)}; total general: {celda(r_total,c_total)}.',
                   f'Filas de residencia: {", ".join(str(r+1) for r in filas)}.',
                   'Extranjero e Ignorado son categorías explícitas de residencia, no departamentos ni subtotales; se conservan sin reinterpretarlas.',
                   f'Se excluyen la fila {matriz[r_total][c_res]!r} y la columna Total del dataset; se conservan en las comprobaciones.',
                   'Notas/fuentes: '+('; '.join(notas) if notas else 'no hay notas ni fuentes dentro de la hoja.'),
                   'Subtotales: ninguno identificado en las filas analíticas.'])
    # Comprobar también cualquier celda fuera del rectángulo identificado.
    extras = [(celda(r,c),v) for r in range(r_total,len(matriz))
              for c,v in enumerate(matriz[r]) if c not in {c_res,c_total,*cols} and v != '']
    if extras:
        errores.append(f'Celdas inesperadas fuera del cuadro: {extras!r}')
    numeros = {}
    for r in [r_total,*filas]:
        for c in [c_total,*cols]:
            try:
                numeros[r,c] = entero(matriz[r][c], celda(r,c))
            except ValueError as error:
                errores.append(str(error))
    lineas.append('Guiones, vacíos o valores no numéricos en el bloque y márgenes: '
                  + str(sum((r,c) not in numeros for r in [r_total,*filas] for c in [c_total,*cols]))
                  + '. No se convierten a cero; cualquiera bloquea la salida.')
    if errores:
        lineas.extend('ERROR: '+e for e in errores)
        return None, None

    ancho = pd.DataFrame([[numeros[r,c] for c in cols] for r in filas], columns=ocurrencias)
    ancho.insert(0,'depto_residencia',residencias)
    largo = ancho.melt(id_vars='depto_residencia',var_name='depto_ocurrencia',value_name='conteo')
    largo.insert(0,'anio',anio)
    largo = largo[COLUMNAS]
    largo['conteo'] = largo['conteo'].astype('int64')
    referencia = {(matriz[r][c_res],matriz[r_enc][c]): numeros[r,c] for r in filas for c in cols}
    reconstruido = {(v.depto_residencia,v.depto_ocurrencia):v.conteo for v in largo.itertuples(index=False)}
    if referencia != reconstruido or largo.duplicated(COLUMNAS[:3]).any():
        errores.append('Intersecciones incorrectas o claves duplicadas tras melt')
    if len(largo) != len(filas)*len(cols):
        errores.append('Cantidad incorrecta de combinaciones')
    lineas.append(f'Intersecciones verificadas individualmente: {len(referencia)}; '
                  f'registros transformados: {len(largo)}.')
    lineas.append('\nVALIDACIÓN POR RESIDENCIA (incluye categorías especiales):')
    por_res = largo.groupby('depto_residencia')['conteo'].sum()
    for r in filas:
        nombre = matriz[r][c_res]
        publicado,calculado = numeros[r,c_total],int(por_res[nombre])
        diferencia = calculado-publicado
        lineas.append(f'{nombre}: publicado={publicado}; transformado={calculado}; diferencia={diferencia}.')
        if diferencia:
            errores.append(f'Total por residencia no coincide: {nombre}')
    lineas.append('\nVALIDACIÓN POR OCURRENCIA:')
    por_oc = largo.groupby('depto_ocurrencia')['conteo'].sum()
    for c in cols:
        nombre = matriz[r_enc][c]
        publicado,calculado = numeros[r_total,c],int(por_oc[nombre])
        diferencia = calculado-publicado
        lineas.append(f'{nombre}: publicado={publicado}; transformado={calculado}; diferencia={diferencia}.')
        if diferencia:
            errores.append(f'Total por ocurrencia no coincide: {nombre}')
    general,calculado = numeros[r_total,c_total],int(largo['conteo'].sum())
    lineas.append(f'\nTOTAL GENERAL: publicado={general}; transformado={calculado}; diferencia={calculado-general}.')
    if calculado != general:
        errores.append('Total general no coincide')
    if sum(numeros[r,c_total] for r in filas) != general or sum(numeros[r_total,c] for c in cols) != general:
        errores.append('Los márgenes originales no coinciden con el total general')
    if hashlib.sha256(ruta.read_bytes()).hexdigest() != hash_original:
        errores.append('El hash del archivo original cambió durante la ejecución')
    lineas.extend(['\nEJEMPLOS DE FILAS TRANSFORMADAS:', largo.head(5).to_csv(index=False).strip()])
    if errores:
        lineas.extend('ERROR: '+e for e in errores)
        lineas.append('Resultado del año: FALLIDO; no se autoriza generar el CSV conjunto.')
        return None,None
    lineas.append('Resultado del año: VALIDADO (intersecciones, filas, columnas, total general y nombres).')
    estructura = dict(dimensiones=dimensiones, fila_ejes=r_ejes+1, fila_encabezados=r_enc+1,
                      fila_total=r_total+1, columna_total=c_total+1,
                      ocurrencias=ocurrencias, residencias=residencias, filas=filas, columnas=cols)
    return largo,estructura


def main():
    lineas = ['VALIDACIÓN IND-002 — SERIE 2009–2024',
              'Transformación con pandas.melt; nombres originales sin normalizar en la salida.',
              'Sin imputaciones: guiones, vacíos, errores y no numéricos bloquean la exportación.',
              'Los ceros numéricos originales sí se conservan.']
    resultados,estructuras = [],{}
    try:
        with CATALOGO.open(encoding='utf-8-sig',newline='') as f:
            indicador = unico([r for r in csv.DictReader(f) if r['id_provisional']=='IND-002'],
                              'Entrada IND-002 del catálogo')
        for anio in ANIOS:
            try:
                nombre = indicador[str(anio)]
                if not nombre or nombre == 'No identificado':
                    raise ValueError('El catálogo no identifica una hoja única')
                largo,estructura = transformar_anio(anio,nombre,lineas)
                if largo is not None:
                    resultados.append(largo); estructuras[anio]=estructura
            except Exception as error:
                lineas.append(f'ERROR {anio}: {type(error).__name__}: {error}')
    except Exception as error:
        lineas.append(f'ERROR de catálogo: {type(error).__name__}: {error}')
    lineas.append('\nCOMPARACIÓN ESTRUCTURAL ENTRE AÑOS CONSECUTIVOS:')
    for anterior, actual in zip(ANIOS, ANIOS[1:]):
        if anterior not in estructuras or actual not in estructuras:
            lineas.append(f'{anterior}→{actual}: comparación no evaluada por fallo en uno de los años.')
            continue
        cambios = [k for k in estructuras[anterior] if estructuras[anterior][k]!=estructuras[actual][k]]
        if cambios:
            for k in cambios:
                lineas.append(f'{anterior}→{actual}, {k}: {estructuras[anterior][k]!r} → {estructuras[actual][k]!r}')
        else:
            lineas.append(f'{anterior}→{actual}: sin diferencias en posiciones, dimensiones, nombres ni orden de los ejes.')
    lineas.append(f'\nAños validados: {len(resultados)}/{len(ANIOS)}.')
    exito = len(resultados)==len(ANIOS)
    if exito:
        datos = pd.concat(resultados,ignore_index=True).sort_values(
            ['anio','depto_residencia','depto_ocurrencia'],kind='stable').reset_index(drop=True)
        SALIDA.parent.mkdir(parents=True,exist_ok=True)
        temporal = SALIDA.with_suffix('.csv.tmp')
        try:
            datos.to_csv(temporal,index=False,encoding='utf-8-sig')
            comprobacion = pd.read_csv(temporal,encoding='utf-8-sig')
            pd.testing.assert_frame_equal(datos,comprobacion)
            temporal.replace(SALIDA)
            lineas.append(f'\nRESULTADO GLOBAL: SERIE VALIDADA. {len(datos)} registros exportados a {SALIDA}.')
            lineas.append('CSV releído y contrastado íntegramente con el DataFrame validado; solo incluye IND-002.')
        except Exception as error:
            exito = False
            lineas.append(f'ERROR al exportar/verificar CSV: {type(error).__name__}: {error}')
        finally:
            temporal.unlink(missing_ok=True)
    if not exito:
        lineas.append('\nRESULTADO GLOBAL: NO VALIDADO. No se generó ni actualizó el CSV conjunto.')
        if SALIDA.exists():
            lineas.append('Existe un CSV de una ejecución anterior; esta ejecución NO lo valida.')
    REPORTE.parent.mkdir(parents=True,exist_ok=True)
    REPORTE.write_text('\n'.join(lineas)+'\n',encoding='utf-8')
    for linea in lineas:
        if linea.startswith(('AÑO:', '\nAÑO:', 'TOTAL GENERAL:', 'Intersecciones verificadas', 'Resultado del año:',
                             'RESULTADO GLOBAL:', 'ERROR', '\nTOTAL GENERAL:', '\nRESULTADO GLOBAL:')):
            print(linea)
    print(f'Reporte: {REPORTE}')
    if not exito:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
